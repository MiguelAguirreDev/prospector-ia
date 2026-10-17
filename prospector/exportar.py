"""Exporta CSV/Excel: listas de emails en lotes (de a 20) o ficha completa."""

from __future__ import annotations

import csv
import re
from pathlib import Path

CABECERA_EMAIL = ["email", "nombre", "telefono", "whatsapp"]
CABECERA_FULL = [
    "nombre",
    "categoria",
    "zona",
    "telefono",
    "whatsapp",
    "email",
    "web",
    "estado_email",
    "fecha_alta",
]


def _escribir_csv(filas: list[list], cabecera: list[str], ruta: Path) -> None:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    with ruta.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(cabecera)
        w.writerows(filas)


def exportar_lote(
    filas: list[list],
    cabecera: list[str],
    carpeta: Path,
    lote: int = 20,
    prefijo: str = "lote",
) -> list[Path]:
    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    generados: list[Path] = []
    for i in range(0, len(filas), lote):
        numero = i // lote + 1
        ruta = carpeta / f"{prefijo}_{numero:02d}.csv"
        _escribir_csv(filas[i : i + lote], cabecera, ruta)
        generados.append(ruta)
    return generados


def exportar_emails(
    db, carpeta: str | Path, lote: int = 20, limite: int | None = None
) -> list[Path]:
    empresas = db.con_email(limite)
    filas = [
        [e["email"], e["nombre"], e["telefono"], e["whatsapp"]]
        for e in empresas
        if e["email"]
    ]
    return exportar_lote(filas, CABECERA_EMAIL, Path(carpeta), lote, prefijo="emails")


def exportar_completo(
    db, ruta_csv: str | Path, limite: int | None = None, solo_objetivo: bool = False
) -> Path:
    empresas = db.listar(limite)
    if solo_objetivo:
        empresas = [e for e in empresas if e.get("es_objetivo")]
    filas = [
        [
            e["nombre"],
            e["categoria"],
            e["zona"],
            e["telefono"],
            e["whatsapp"],
            e["email"],
            e["web"],
            e["estado_email"],
            e["fecha_alta"],
        ]
        for e in empresas
    ]
    ruta_csv = Path(ruta_csv)
    _escribir_csv(filas, CABECERA_FULL, ruta_csv)
    _xlsx(ruta_csv.with_suffix(".xlsx"), filas, CABECERA_FULL)
    return ruta_csv


def nombre_archivo(nombre: str) -> str:
    """Convierte un nombre de lista en algo seguro para usar como archivo."""
    limpio = re.sub(r'[\\/:*?"<>|]+', "_", (nombre or "").strip())
    limpio = re.sub(r"\s+", " ", limpio).strip(" .")
    return limpio[:80] or "lista"


def exportar_lista(db, nombre: str, carpeta: str | Path) -> Path:
    """Escribe <carpeta>/<nombre>.csv y .xlsx con las empresas de esa lista."""
    empresas = db.empresas_de_lista(nombre)
    filas = [
        [
            e["nombre"],
            e["categoria"],
            e["zona"],
            e["telefono"],
            e["whatsapp"],
            e["email"],
            e["web"],
            e["estado_email"],
            e["fecha_alta"],
        ]
        for e in empresas
    ]
    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    ruta = carpeta / f"{nombre_archivo(nombre)}.csv"
    _escribir_csv(filas, CABECERA_FULL, ruta)
    _xlsx(ruta.with_suffix(".xlsx"), filas, CABECERA_FULL)
    return ruta


def exportar_listas(db, carpeta: str | Path) -> list[Path]:
    """Escribe un Excel por cada lista que al menos tenga una empresa."""
    rutas: list[Path] = []
    for l in db.listar_listas():
        if l.get("n"):
            rutas.append(exportar_lista(db, l["nombre"], carpeta))
    return rutas


def _xlsx(ruta: Path, filas: list[list], cabecera: list[str]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
    except Exception:
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Empresas"
    ws.append(cabecera)
    for fila in filas:
        ws.append(fila)
    for i, _ in enumerate(cabecera, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 24
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(ruta))
